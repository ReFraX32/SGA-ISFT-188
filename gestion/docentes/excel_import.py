# --------------------------------------------------------------------------
# MÓDULO: Gestión de Docentes — IMPORTACIÓN DESDE EXCEL
# --------------------------------------------------------------------------
# Reutiliza las funciones de normalización (normalizar_dni, normalizar_cuil,
# etc.) que ya existen en gestion/excel_import_export.py para el importador
# de alumnos, así ambos importadores quedan consistentes y no hay lógica
# de limpieza de datos duplicada en dos lugares distintos del proyecto.
# --------------------------------------------------------------------------

import openpyxl

from gestion.models import Persona, Docente
from gestion.excel_import_export import (
    normalizar_dni,
    normalizar_cuil,
    normalizar_nombre_propio,
    parsear_fecha_flexible,
    normalizar_genero,
    normalizar_mail,
)


def procesar_importacion_docentes_excel(file_obj):
    """
    Procesa un archivo Excel subido por el usuario para registrar o actualizar docentes.
    Normaliza datos de Persona y crea/actualiza el perfil Docente asociado.
    Retorna un diccionario summary con el resultado de la operación.
    """
    summary = {
        'success': False,
        'creados': 0,
        'actualizados': 0,
        'filas_procesadas': 0,
        'errores': [],
        'advertencias': [],
        'mensaje': ''
    }

    try:
        wb = openpyxl.load_workbook(file_obj, data_only=True)
    except Exception as e:
        summary['mensaje'] = f"El archivo proporcionado no es un archivo Excel (.xlsx) válido: {str(e)}"
        return summary

    ws = None
    for sheet_name in ["Carga de Docentes", "Docentes", "Carga de Profesores", "Profesores", "Hoja 1", "Sheet1"]:
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            break
    if not ws:
        ws = wb.active

    if not ws or ws.max_row < 2:
        summary['mensaje'] = "La hoja de cálculo no contiene filas suficientes de datos para procesar."
        return summary

    # Detectar fila de encabezados buscando la fila que contenga DNI, Apellido y Nombre
    header_row = None
    col_map = {}
    for r in range(1, min(10, ws.max_row + 1)):
        row_vals = [str(ws.cell(r, c).value or "").strip().upper() for c in range(1, ws.max_column + 1)]
        temp_map = {}
        for c_idx, val in enumerate(row_vals, 1):
            if "DNI" in val or "DOCUMENTO" in val: temp_map["dni"] = c_idx
            elif "APELLIDO" in val: temp_map["apellido"] = c_idx
            elif "NOMBRE" in val: temp_map["nombre"] = c_idx
            elif "TITULO" in val or "TÍTULO" in val or "MATRICULA" in val or "MATRÍCULA" in val: temp_map["titulo"] = c_idx
            elif "CUIL" in val: temp_map["cuil"] = c_idx
            elif "FECHA" in val or "NACIMIENTO" in val: temp_map["fecha_nac"] = c_idx
            elif "GENERO" in val or "GÉNERO" in val or "IDENTIDAD" in val: temp_map["genero"] = c_idx
            elif "NACIONALIDAD" in val: temp_map["nacionalidad"] = c_idx
            elif "LOCALIDAD" in val or "CIUDAD" in val: temp_map["localidad"] = c_idx
            elif "DOMICILIO" in val or "DIRECCION" in val or "DIRECCIÓN" in val: temp_map["domicilio"] = c_idx
            elif "TELEFONO" in val or "TELÉFONO" in val or "CELULAR" in val: temp_map["telefono"] = c_idx
            elif "MAIL" in val or "CORREO" in val or "EMAIL" in val: temp_map["mail"] = c_idx

        if "dni" in temp_map and "apellido" in temp_map and "nombre" in temp_map:
            header_row = r
            col_map = temp_map
            break

    if not header_row or "dni" not in col_map or "apellido" not in col_map or "nombre" not in col_map:
        summary['mensaje'] = "No se pudieron identificar las columnas obligatorias (DNI, Apellido, Nombre) en el archivo Excel. Utilice la plantilla oficial de docentes."
        return summary

    # Procesar filas de datos
    for r in range(header_row + 1, ws.max_row + 1):
        dni_raw = ws.cell(r, col_map["dni"]).value
        apellido_raw = ws.cell(r, col_map["apellido"]).value
        nombre_raw = ws.cell(r, col_map["nombre"]).value

        # Si toda la fila obligatoria está vacía, continuar
        if not dni_raw and not apellido_raw and not nombre_raw:
            continue

        dni_clean = normalizar_dni(dni_raw)
        if not dni_clean:
            summary['errores'].append(f"Fila {r}: DNI inválido o ausente ('{dni_raw}').")
            continue

        apellido_clean = normalizar_nombre_propio(apellido_raw)
        nombre_clean = normalizar_nombre_propio(nombre_raw)

        if not apellido_clean or not nombre_clean:
            summary['errores'].append(f"Fila {r} (DNI {dni_clean}): Falta completar el Apellido o Nombre del docente.")
            continue

        # Campos opcionales normalizados
        titulo_raw = ws.cell(r, col_map["titulo"]).value if "titulo" in col_map else None
        titulo_clean = str(titulo_raw).strip() if (titulo_raw and str(titulo_raw).strip() not in ['', 'None', '-']) else None

        cuil_raw = ws.cell(r, col_map["cuil"]).value if "cuil" in col_map else None
        cuil_clean = normalizar_cuil(cuil_raw) or None

        fecha_nac_raw = ws.cell(r, col_map["fecha_nac"]).value if "fecha_nac" in col_map else None
        fecha_nac_clean = parsear_fecha_flexible(fecha_nac_raw)

        genero_raw = ws.cell(r, col_map["genero"]).value if "genero" in col_map else None
        genero_clean = normalizar_genero(genero_raw)

        nacionalidad_raw = ws.cell(r, col_map["nacionalidad"]).value if "nacionalidad" in col_map else None
        nacionalidad_clean = normalizar_nombre_propio(nacionalidad_raw) if (nacionalidad_raw and str(nacionalidad_raw).strip() not in ['', 'None', '-']) else "Argentina"

        localidad_raw = ws.cell(r, col_map["localidad"]).value if "localidad" in col_map else None
        localidad_clean = normalizar_nombre_propio(localidad_raw) if (localidad_raw and str(localidad_raw).strip() not in ['', 'None', '-']) else None

        domicilio_raw = ws.cell(r, col_map["domicilio"]).value if "domicilio" in col_map else None
        domicilio_clean = str(domicilio_raw).strip() if domicilio_raw else None

        telefono_raw = ws.cell(r, col_map["telefono"]).value if "telefono" in col_map else None
        telefono_clean = str(telefono_raw).strip() if telefono_raw else None

        mail_raw = ws.cell(r, col_map["mail"]).value if "mail" in col_map else None
        mail_clean = normalizar_mail(mail_raw)

        # Crear o actualizar Persona
        persona_obj, created_persona = Persona.objects.update_or_create(
            dni=dni_clean,
            defaults={
                'apellido': apellido_clean,
                'nombre': nombre_clean,
                'cuil': cuil_clean,
                'fecha_nacimiento': fecha_nac_clean,
                'identidad': genero_clean,
                'nacionalidad': nacionalidad_clean,
                'localidad': localidad_clean,
                'domicilio': domicilio_clean,
                'telefono': telefono_clean,
                'mail': mail_clean
            }
        )

        # Crear o actualizar Docente
        docente_defaults = {}
        if titulo_clean is not None:
            docente_defaults['titulo_mn'] = titulo_clean

        docente_obj, created_docente = Docente.objects.update_or_create(
            persona=persona_obj,
            defaults=docente_defaults
        )

        if created_docente:
            summary['creados'] += 1
        else:
            summary['actualizados'] += 1

        summary['filas_procesadas'] += 1

    summary['success'] = True
    summary['mensaje'] = f"Importación finalizada con éxito: {summary['creados']} docentes nuevos registrados y {summary['actualizados']} docentes actualizados."
    return summary
